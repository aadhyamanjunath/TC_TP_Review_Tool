#!/usr/bin/env python3
"""logic.py

Core logic for the Stack/Overflow checkpoint review tool.

Split out from `checkpoint_new.py`.
- `logic.py` contains all business logic (parsing, reviewing, report generation)
- `main.py` contains only CLI argument parsing and orchestration
- `__init__.py` exposes a small public API for programmatic use
"""

from __future__ import annotations

import html
import re
from collections import defaultdict
from pathlib import Path

import pandas as pd

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

# -----------------------------------------------------------------------------
# Defaults (your folder structure)
# -----------------------------------------------------------------------------

DEFAULT_BASE = Path.home() / "Documents" / "checkpoint_review_tool"
DEFAULT_INPUT_DIR = DEFAULT_BASE / "input_docs"
DEFAULT_OUTPUT_DIR = DEFAULT_BASE / "output_docs"
DEFAULT_SRS = DEFAULT_INPUT_DIR / "SCU_Full_SRS.docx"
DEFAULT_TRACE = DEFAULT_INPUT_DIR / "Traceability.xlsx"
DEFAULT_TC = DEFAULT_INPUT_DIR / "MRJ_SCU_STC_SRS_1162_POC_Test_Case.xlsx"
DEFAULT_OUT = DEFAULT_OUTPUT_DIR / "Stack_Checkpoint_Review_Report.xlsx"
DEFAULT_KEYWORDS = "stack overflow;stack overrun;stack usage;stack margin;stack;overflow;overrun"
DEFAULT_REQ_ID_REGEX = r"\b[A-Z]{2,}_[A-Z0-9]+(?:_[A-Z0-9]+)*_\d+\b"


# -----------------------------------------------------------------------------
# Utility helpers
# -----------------------------------------------------------------------------

def ensure_exists(path: Path, label: str) -> None:
    if not path.exists():
        raise FileNotFoundError(f"{label} file not found: {path}")


def normalize(s: str) -> str:
    s = html.unescape(str(s or ""))
    return re.sub(r"\s+", " ", s.strip()).lower()


def safe_str(x) -> str:
    if x is None:
        return ""
    if isinstance(x, float) and pd.isna(x):
        return ""
    return str(x).strip()


def safe_sheet(sheet_name):
    """Avoid pandas returning dict when sheet_name=None."""
    return 0 if sheet_name in (None, "", "None") else sheet_name


def find_best_column(df_cols, candidates):
    cols = list(df_cols)
    norm_cols = {normalize(c): c for c in cols}

    # exact
    for cand in candidates:
        key = normalize(cand)
        if key in norm_cols:
            return norm_cols[key]

    # contains
    for c in cols:
        nc = normalize(c)
        for cand in candidates:
            if normalize(cand) in nc:
                return c
    return None


def to_float_percent(x):
    """Convert 70, 70.0, "70", "70%", " 70 % " to float(70.0). Return None if not possible."""
    s = safe_str(x)
    if not s:
        return None
    s = s.replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def extract_threshold_from_text(text: str):
    """Extract threshold like 70% from requirement text."""
    t = html.unescape(text or "")
    m = re.search(r"(\d{1,3})\s*%", t)
    return float(m.group(1)) if m else None


# -----------------------------------------------------------------------------
# Step 1: Extract requirements from SRS DOCX
# -----------------------------------------------------------------------------

def extract_requirements_from_docx(srs_path: Path, req_id_regex: str) -> pd.DataFrame:
    if Document is None:
        raise RuntimeError("python-docx not installed. Install: pip install python-docx")

    doc = Document(str(srs_path))
    req_id_re = re.compile(req_id_regex)
    found = []

    # tables
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if not cells:
                continue
            m = req_id_re.search(cells[0] or "")
            if m:
                rid = m.group(0)
                rtxt = cells[1] if len(cells) > 1 else ""
                found.append((rid, rtxt))

    # paragraphs fallback
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        m = req_id_re.search(t)
        if m:
            rid = m.group(0)
            found.append((rid, t))

    # dedupe keep longest per ID
    best = {}
    for rid, txt in found:
        if rid not in best or len(txt or "") > len(best[rid] or ""):
            best[rid] = txt

    df = (
        pd.DataFrame(
            [{"Requirement ID": k, "Requirement Text": v} for k, v in best.items()]
        )
        .sort_values("Requirement ID")
        .reset_index(drop=True)
    )
    return df


def filter_requirements_by_keywords(req_df: pd.DataFrame, keywords):
    kw_norm = [normalize(k) for k in keywords if k.strip()]

    def hit_list(txt):
        t = normalize(txt)
        hits = [k for k in kw_norm if k and k in t]
        return ", ".join(sorted(set(hits))) if hits else ""

    out = req_df.copy()
    out["Keyword Hits"] = out["Requirement Text"].apply(hit_list)
    out["Is Relevant?"] = out["Keyword Hits"].apply(lambda x: "YES" if x else "NO")
    out["Threshold (%) (from SRS)"] = out["Requirement Text"].apply(
        extract_threshold_from_text
    )
    return out


# -----------------------------------------------------------------------------
# Step 2: Load traceability mapping
# -----------------------------------------------------------------------------

def load_traceability_map(trace_path: Path, req_id_regex: str, sheet_name=None):
    """Load ReqID -> [TC_IDs] mapping.

    Supports:
    - long format: [Requirement ID, Test Case ID]
    - summary format: [Requirement ID, Mapped Test Cases (comma-separated)]
    """

    req_id_re = re.compile(req_id_regex)
    df = pd.read_excel(trace_path, engine="openpyxl", sheet_name=safe_sheet(sheet_name))

    req_col = find_best_column(df.columns, ["Requirement ID", "Req ID", "Requirement", "SRS ID"])
    tc_col = find_best_column(df.columns, ["Test Case ID", "TC ID", "Testcase ID", "Test Case"])

    if not tc_col:
        tc_col = find_best_column(
            df.columns,
            ["Mapped Test Cases", "Mapped Test Cases (comma-separated)", "Mapped TCs"],
        )

    if req_col and not tc_col and len(df.columns) == 2:
        tc_col = [c for c in df.columns if c != req_col][0]

    if not req_col or not tc_col:
        raise ValueError(
            "Could not detect traceability columns.\n"
            f"Found columns: {list(df.columns)}\n"
            "Expected 'Requirement ID' and a TC column."
        )

    mapping = defaultdict(list)
    for _, r in df.iterrows():
        rid_raw = safe_str(r.get(req_col, ""))
        tc_raw = safe_str(r.get(tc_col, ""))
        m = req_id_re.search(rid_raw)
        if not m:
            continue
        rid = m.group(0)

        tc_list = re.split(r"[,\n;/]+", tc_raw)
        for tcid in tc_list:
            tcid = tcid.strip()
            if tcid and tcid.lower() != "nan":
                mapping[rid].append(tcid)

    for rid in list(mapping.keys()):
        mapping[rid] = sorted(set(mapping[rid]))

    return mapping


# -----------------------------------------------------------------------------
# Step 3: Load test case scenarios sheet
# -----------------------------------------------------------------------------

def load_testcases(tc_path: Path, sheet_name=None):
    df = pd.read_excel(tc_path, engine="openpyxl", sheet_name=safe_sheet(sheet_name))

    required = [
        "Requirement_ID",
        "TC_ID",
        "Scenario_No",
        "Title",
        "Test Case Purpose/Objective",
        "Input Format (POC)",
        "Threshold (%)",
        "Input Dataset / Values",
        "Expected Stack Space Usage %",
        "Expected Message output",
        "Notes",
    ]

    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(
            "TC file does not match expected columns.\n"
            f"Missing columns: {missing}\n"
            f"Found columns: {list(df.columns)}"
        )

    return df


# -----------------------------------------------------------------------------
# Parsing input dataset values to compute stack usage
# -----------------------------------------------------------------------------

def parse_used_alloc_pairs(dataset_text: str):
    """Extract (used, allocated) pairs from 'Input Dataset / Values'."""

    text = safe_str(dataset_text)
    if not text:
        return [], ["Empty dataset"]

    invalid = []
    pairs = []

    if re.search(r"\bused\s*=\s*null\b", text, flags=re.IGNORECASE):
        invalid.append("used is NULL")

    kv_pattern = re.compile(
        r"(?:alloc(?:ated)?\s*=\s*(?P<alloc>-?\d+(?:\.\d+)?))"
        r".{0,80}?"
        r"(?:used\s*=\s*(?P<used>-?\d+(?:\.\d+)?))",
        flags=re.IGNORECASE,
    )

    for m in kv_pattern.finditer(text):
        alloc = float(m.group("alloc"))
        used = float(m.group("used"))
        pairs.append((used, alloc))

    frac_pattern = re.compile(
        r"(?P<used>-?\d+(?:\.\d+)?)\s*/\s*(?P<alloc>-?\d+(?:\.\d+)?)"
    )

    for m in frac_pattern.finditer(text):
        used = float(m.group("used"))
        alloc = float(m.group("alloc"))
        pairs.append((used, alloc))

    if not pairs:
        if re.search(r"used\s*=\s*'[^']+'", text, flags=re.IGNORECASE):
            invalid.append("used is non-numeric string")
        if re.search(r"alloc(?:ated)?\s*=\s*'[^']+'", text, flags=re.IGNORECASE):
            invalid.append("allocated is non-numeric/unit string")
        invalid.append("No usable (used,allocated) numeric pairs parsed")

    valid_pairs = []
    for used, alloc in pairs:
        if alloc <= 0:
            invalid.append(f"allocated <= 0 (alloc={alloc})")
            continue
        if used < 0:
            invalid.append(f"used < 0 (used={used})")
            continue
        valid_pairs.append((used, alloc))

    return valid_pairs, sorted(set(invalid))


def compute_max_usage_pct(pairs):
    if not pairs:
        return None
    return max((used / alloc) * 100.0 for used, alloc in pairs)


# -----------------------------------------------------------------------------
# Reviewer logic
# -----------------------------------------------------------------------------

def expected_category_from_usage(max_usage, threshold):
    if max_usage is None:
        return "INVALID_INPUT"
    if threshold is None:
        return "UNKNOWN"
    return "ERROR" if max_usage > threshold else "OK"


def message_category_from_text(msg: str):
    m = normalize(msg)
    if not m:
        return "MISSING"
    if "invalid_input" in m or "invalid input" in m or "cannot compute" in m:
        return "INVALID_INPUT"
    if "error" in m or "exceed" in m or "exceeded" in m:
        return "ERROR"
    if "ok" in m or "within limit" in m or "within" in m:
        return "OK"
    return "OTHER"


def approx_equal(a, b, tol=0.35):
    if a is None or b is None:
        return False
    return abs(a - b) <= tol


def review_scenario_row(row: pd.Series, requirement_text: str):
    threshold_tc = to_float_percent(row.get("Threshold (%)"))
    threshold_srs = extract_threshold_from_text(requirement_text)
    threshold_used = threshold_tc if threshold_tc is not None else threshold_srs

    dataset = safe_str(row.get("Input Dataset / Values"))
    pairs, invalid_reasons = parse_used_alloc_pairs(dataset)
    max_usage = compute_max_usage_pct(pairs)

    expected_usage_num = to_float_percent(row.get("Expected Stack Space Usage %"))
    expected_msg = safe_str(row.get("Expected Message output"))

    expected_msg_cat = message_category_from_text(expected_msg)
    tool_expected_cat = expected_category_from_usage(max_usage, threshold_used)

    fail_reasons = []

    if max_usage is None and expected_msg_cat != "INVALID_INPUT":
        fail_reasons.append(
            f"Dataset invalid ({'; '.join(invalid_reasons)}), but expected message category is '{expected_msg_cat}'"
        )

    if max_usage is not None and expected_usage_num is not None:
        if not approx_equal(expected_usage_num, max_usage):
            fail_reasons.append(
                f"Expected usage% ({expected_usage_num}) != parsed max usage% ({max_usage:.3f})"
            )

    if tool_expected_cat != "UNKNOWN" and expected_msg_cat != tool_expected_cat:
        fail_reasons.append(
            "Expected message implies '{}' but rule expects '{}' "
            "(max_usage={}, threshold={})".format(
                expected_msg_cat,
                tool_expected_cat,
                None if max_usage is None else round(max_usage, 3),
                threshold_used,
            )
        )

    verdict = "PASS" if not fail_reasons else "FAIL"

    return {
        "Threshold (%) used": threshold_used,
        "Parsed Max Stack Usage %": None if max_usage is None else round(max_usage, 6),
        "Expected Usage % (numeric)": expected_usage_num,
        "Expected Message Category": expected_msg_cat,
        "Tool-Expected Category": tool_expected_cat,
        "Verdict": verdict,
        "Fail Reason": "OK" if verdict == "PASS" else "\n".join(fail_reasons),
    }


# -----------------------------------------------------------------------------
# Build output report
# -----------------------------------------------------------------------------

def build_report(req_df_filtered, trace_map, tc_df, out_path: Path):
    req_text_map = dict(
        zip(req_df_filtered["Requirement ID"], req_df_filtered["Requirement Text"])
    )

    relevant_reqs = req_df_filtered[req_df_filtered["Is Relevant?"] == "YES"].copy()
    relevant_ids = relevant_reqs["Requirement ID"].tolist()

    rows = []
    missing_trace = []

    for rid in relevant_ids:
        linked_tc_ids = trace_map.get(rid, [])
        if not linked_tc_ids:
            missing_trace.append(rid)
            continue

        tc_subset = tc_df[tc_df["TC_ID"].astype(str).isin(set(linked_tc_ids))].copy()

        if tc_subset.empty:
            rows.append(
                {
                    "Requirement ID": rid,
                    "Test Case ID": "",
                    "Scenario No": "",
                    "Title": "",
                    "Input Dataset / Values": "",
                    "Threshold (%) (TC)": "",
                    "Expected Stack Space Usage %": "",
                    "Expected Message output": "",
                    "Verdict": "FAIL",
                    "Fail Reason": "Traceability links TCs, but no matching TC rows found in TC file.",
                }
            )
            continue

        req_text = req_text_map.get(rid, "")

        for _, r in tc_subset.iterrows():
            review = review_scenario_row(r, req_text)

            input_section = "\n".join(
                [
                    f"Scenario: {safe_str(r.get('Scenario_No'))}",
                    f"Title: {safe_str(r.get('Title'))}",
                    f"Objective: {safe_str(r.get('Test Case Purpose/Objective'))}",
                    f"Input Format (POC): {safe_str(r.get('Input Format (POC)'))}",
                    f"Notes: {safe_str(r.get('Notes'))}",
                ]
            ).strip()

            rows.append(
                {
                    "Requirement ID": safe_str(r.get("Requirement_ID")) or rid,
                    "Requirement Text": req_text,
                    "Test Case ID": safe_str(r.get("TC_ID")),
                    "Scenario No": safe_str(r.get("Scenario_No")),
                    "Title": safe_str(r.get("Title")),
                    "Input section": input_section,
                    "Input Dataset / Values": safe_str(r.get("Input Dataset / Values")),
                    "Threshold (%) (TC)": safe_str(r.get("Threshold (%)")),
                    "Threshold (%) used": review["Threshold (%) used"],
                    "Parsed Max Stack Usage %": review["Parsed Max Stack Usage %"],
                    "Expected Stack Space Usage %": safe_str(
                        r.get("Expected Stack Space Usage %")
                    ),
                    "Expected Usage % (numeric)": review["Expected Usage % (numeric)"],
                    "Expected Message output": safe_str(r.get("Expected Message output")),
                    "Expected Message Category": review["Expected Message Category"],
                    "Tool-Expected Category": review["Tool-Expected Category"],
                    "Verdict": review["Verdict"],
                    "Fail Reason": review["Fail Reason"],
                }
            )

    review_df = pd.DataFrame(rows)

    if review_df.empty:
        summary_df = pd.DataFrame(
            [
                {
                    "Total Scenarios": 0,
                    "PASS": 0,
                    "FAIL": 0,
                    "Requirements w/o Trace Links": ", ".join(missing_trace)
                    if missing_trace
                    else "None",
                    "Notes": "No relevant stack requirements or scenario rows found.",
                }
            ]
        )
    else:
        summary_df = pd.DataFrame(
            [
                {
                    "Total Scenarios": int(len(review_df)),
                    "PASS": int((review_df["Verdict"] == "PASS").sum()),
                    "FAIL": int((review_df["Verdict"] == "FAIL").sum()),
                    "Requirements w/o Trace Links": ", ".join(missing_trace)
                    if missing_trace
                    else "None",
                }
            ]
        )

    if not review_df.empty:
        by_req = (
            review_df.groupby("Requirement ID")["Verdict"]
            .value_counts()
            .unstack(fill_value=0)
            .reset_index()
        )
    else:
        by_req = pd.DataFrame(columns=["Requirement ID", "PASS", "FAIL"])

    out_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        review_df.to_excel(writer, sheet_name="Scenario_Review", index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        by_req.to_excel(writer, sheet_name="By_Requirement", index=False)
        relevant_reqs.to_excel(writer, sheet_name="Stack_Requirements", index=False)

    print(f"\n✅ Report generated successfully: {out_path}")


# -----------------------------------------------------------------------------
# High-level runner (usable from CLI or programmatically)
# -----------------------------------------------------------------------------

def run(
    srs: Path = DEFAULT_SRS,
    trace: Path = DEFAULT_TRACE,
    tc: Path = DEFAULT_TC,
    out: Path = DEFAULT_OUT,
    *,
    keywords: str = DEFAULT_KEYWORDS,
    req_id_regex: str = DEFAULT_REQ_ID_REGEX,
    trace_sheet=None,
    tc_sheet="Test Cases",
) -> Path:
    """Run the full pipeline and produce the output report."""

    srs_path = Path(srs)
    trace_path = Path(trace)
    tc_path = Path(tc)
    out_path = Path(out)

    ensure_exists(srs_path, "SRS")
    ensure_exists(trace_path, "Traceability")
    ensure_exists(tc_path, "Test Cases")

    keyword_list = [k.strip() for k in keywords.split(";") if k.strip()]

    req_df = extract_requirements_from_docx(srs_path, req_id_regex)
    if req_df.empty:
        raise RuntimeError("No requirements found in SRS. Check regex or SRS formatting.")

    req_df_filtered = filter_requirements_by_keywords(req_df, keyword_list)
    trace_map = load_traceability_map(trace_path, req_id_regex, sheet_name=trace_sheet)
    tc_df = load_testcases(tc_path, sheet_name=tc_sheet)

    build_report(req_df_filtered, trace_map, tc_df, out_path)
    return out_path

