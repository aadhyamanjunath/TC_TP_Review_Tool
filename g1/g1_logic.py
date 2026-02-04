#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""DO-178 Test Case Review Tool - core logic

This module contains the implementation for:
- Empty-cell detection
- DOCX text loading
- Auto-numbered heading number extraction from DOCX XML
- Section reference parsing and verification
- Instruction detection
- Equipment extraction and verification against DOCX
- Testcase file loading (Excel/CSV)
- Review logic + report writing

The CLI is implemented in main.py.
"""

import argparse
import os
import re
import sys
import time
import tempfile
import shutil
from typing import List, Optional, Set, Tuple

import zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict

import pandas as pd
from docx import Document


# ---------------------------
# Robust empty detection
# ---------------------------
EMPTY_LIKE = {"", "na", "n/a", "none", "null", "-", "--"}

def is_empty_cell(val) -> bool:
    """True if NaN/None, whitespace-only (incl NBSP), or placeholder like N/A, '-'."""
    if val is None or pd.isna(val):
        return True
    s = str(val)
    s = s.replace("\u00A0", " ")         # NBSP -> space
    s = re.sub(r"\s+", " ", s).strip()   # normalize whitespace
    return (s.lower() in EMPTY_LIKE) or (s == "")


# ---------------------------
# Load .docx as searchable text
# ---------------------------
def load_docx_text(docx_path: str) -> str:
    """Extract visible text from doc body paragraphs and tables (lowercased)."""
    if not docx_path.lower().endswith(".docx"):
        raise ValueError("Only .docx supported. If you have .doc, open in Word and Save As .docx.")
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"Reference document not found: {docx_path}")

    doc = Document(docx_path)
    parts = []

    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    parts.append(t)

    text = " ".join(parts)
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


# ---------------------------
# Extract auto-numbered section labels (e.g., 3.2.1) from Word DOCX XML
# ---------------------------
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

def _qn(tag: str) -> str:
    """Qualified name helper for ElementTree."""
    prefix, local = tag.split(":")
    return f"{{{NS[prefix]}}}{local}"

def extract_section_numbers_from_docx(docx_path: str) -> set[str]:
    """
    Extract auto-numbered labels like '3.2.1' from a Word document by interpreting
    numbering.xml + document.xml.

    Returns: set of section-number strings (e.g., {'1', '1.1', '3.2.1', ...})
    """
    with zipfile.ZipFile(docx_path) as z:
        doc_xml = z.read("word/document.xml")
        numbering_xml = z.read("word/numbering.xml") if "word/numbering.xml" in z.namelist() else None

    if not numbering_xml:
        return set()

    doc_tree = ET.fromstring(doc_xml)
    num_tree = ET.fromstring(numbering_xml)

    # numId -> abstractNumId
    numId_to_abs = {}
    for num in num_tree.findall(".//w:num", NS):
        numId = num.get(_qn("w:numId"))
        abs_el = num.find("./w:abstractNumId", NS)
        if numId and abs_el is not None:
            abs_id = abs_el.get(_qn("w:val"))
            if abs_id is not None:
                numId_to_abs[numId] = abs_id

    # abstractNumId -> lvlText per ilvl (e.g., "%1.%2.%3")
    abs_to_lvltext = defaultdict(dict)
    for absnum in num_tree.findall(".//w:abstractNum", NS):
        abs_id = absnum.get(_qn("w:abstractNumId"))
        if abs_id is None:
            continue
        for lvl in absnum.findall("./w:lvl", NS):
            ilvl = lvl.get(_qn("w:ilvl"))
            lvlText_el = lvl.find("./w:lvlText", NS)
            if ilvl is not None and lvlText_el is not None:
                fmt = lvlText_el.get(_qn("w:val"))
                if fmt:
                    abs_to_lvltext[abs_id][int(ilvl)] = fmt

    counters = defaultdict(lambda: defaultdict(int))
    found_numbers = set()

    for p in doc_tree.findall(".//w:p", NS):
        pPr = p.find("./w:pPr", NS)
        if pPr is None:
            continue
        numPr = pPr.find("./w:numPr", NS)
        if numPr is None:
            continue

        numId_el = numPr.find("./w:numId", NS)
        ilvl_el = numPr.find("./w:ilvl", NS)
        if numId_el is None or ilvl_el is None:
            continue

        numId = numId_el.get(_qn("w:val"))
        ilvl = ilvl_el.get(_qn("w:val"))
        if numId is None or ilvl is None:
            continue

        try:
            level = int(ilvl)
        except ValueError:
            continue

        abs_id = numId_to_abs.get(numId)
        if abs_id is None:
            continue

        # Increment current level, reset deeper levels
        counters[numId][level] += 1
        for dl in list(counters[numId].keys()):
            if dl > level:
                counters[numId][dl] = 0

        lvlText = abs_to_lvltext.get(abs_id, {}).get(level)
        if lvlText:
            label = lvlText
            for k in range(1, 10):  # levels 1..9
                cnt = counters[numId].get(k - 1, 0)
                label = label.replace(f"%{k}", str(cnt))
            label = label.strip()
            label = re.sub(r"[)\].:;\-]+$", "", label).strip()
        else:
            parts = [str(counters[numId].get(lvl, 0)) for lvl in range(level + 1)]
            label = ".".join([p for p in parts if p != "0"]).strip()

        if re.fullmatch(r"\d+(?:\.\d+)*", label):
            found_numbers.add(label)

    return found_numbers


# ---------------------------
# Section reference extraction and verification
# ---------------------------
SECTION_REF_PATTERNS = [
    r"(?:refer(?:\s+the)?|ref(?:erence)?|see|as per|per)\s+(?:to\s+)?(?:section|sec\.?)\s*([0-9]+(?:\.[0-9]+)*)",
    r"(?:section|sec\.?)\s*([0-9]+(?:\.[0-9]+)*)",
]

def extract_section_refs(text: str) -> List[str]:
    if not isinstance(text, str) or not text.strip():
        return []
    found = []
    for pat in SECTION_REF_PATTERNS:
        for m in re.finditer(pat, text, flags=re.IGNORECASE):
            num = m.group(1).strip(". ")
            if num and num not in found:
                found.append(num)
    return found

def section_exists_in_doc(section_num: str, doc_text: str, section_numbers: set[str]) -> bool:
    """
    Works for BOTH:
    - auto-numbered headings (checked via section_numbers set)
    - literal text references (fallback substring search)
    """
    s = str(section_num).strip()
    s = s.replace("Section", "").replace("section", "").replace("sec.", "").replace("Sec.", "")
    s = re.sub(r"\s+", "", s)

    if s in section_numbers:
        return True

    d = doc_text.lower()
    return (s in d) or (f"section {s}" in d) or (f"sec. {s}" in d)


# ---------------------------
# Instruction detection
# ---------------------------
ACTION_VERBS = [
    "click", "press", "enter", "select", "choose", "verify", "confirm", "ensure",
    "navigate", "open", "close", "type", "input", "check", "compare", "record",
    "observe", "monitor", "power on", "power off", "connect", "disconnect",
    "install", "uninstall", "run", "execute", "start", "stop", "set", "reset",
    "configure", "measure", "toggle", "enable", "disable", "apply", "note"
]

STEP_PATTERNS = [
    r"^\s*\d+\s*[\)\.\:\-]\s+",
    r"^\s*\(\d+\)\s+",
    r"^\s*[a-zA-Z]\s*[\)\.\:]\s+",
    r"^\s*[\-\*•]\s+",
    r"\bstep\s*\d+\b",
]

def looks_like_instructions(text: str) -> bool:
    if not isinstance(text, str):
        return False
    t = text.strip()
    if not t:
        return False

    lines = [ln.strip() for ln in re.split(r"[\r\n]+", t) if ln.strip()]
    for ln in lines:
        for pat in STEP_PATTERNS:
            if re.search(pat, ln, flags=re.IGNORECASE):
                return True

    low = " " + re.sub(r"\s+", " ", t.lower()) + " "
    for v in ACTION_VERBS:
        if re.search(rf"\b{re.escape(v)}\b", low):
            return True

    return False


# ---------------------------
# Equipment detection + "must exist in docx" verification
# ---------------------------
DEFAULT_EQUIPMENT_KEYWORDS = [
    "hil", "hil bench", "hil rig", "rig", "test bench", "simulator", "stimulator",
    "psu", "power supply", "oscilloscope", "scope", "multimeter", "dmm",
    "logic analyzer", "spectrum analyzer", "function generator", "signal generator",
    "daq", "ni-daq", "pxi", "pxie", "cdaq",
    "arinc", "arinc-429", "afdx", "can", "canoe", "lin", "flexray", "mil-std-1553",
    "rs-232", "rs-422", "ethernet", "udp", "tcp", "serial",
    "ecu", "lru", "sru", "jtag", "usb", "sd card", "sd-card",
    "load bank", "environmental chamber", "thermal chamber"
]

MODEL_TOKEN_PATTERNS = [
    r"\b[A-Z]{2,}[A-Z0-9\-]*\d+[A-Z0-9\-]*\b",
    r"\b[A-Z]+\-\d+[A-Z0-9\-]*\b",
    r"\b\d{2,}[A-Z][A-Z0-9\-]*\b",
]

def load_equipment_keywords(path: Optional[str]) -> List[str]:
    if not path:
        return DEFAULT_EQUIPMENT_KEYWORDS
    if not os.path.exists(path):
        raise FileNotFoundError(f"Equipment list file not found: {path}")
    txt = open(path, "r", encoding="utf-8", errors="ignore").read()
    toks = []
    for piece in re.split(r"[\r\n,;]+", txt):
        piece = piece.strip()
        if piece:
            toks.append(piece)
    return toks or DEFAULT_EQUIPMENT_KEYWORDS

def extract_equipment(text: str, keywords: List[str]) -> List[str]:
    """Extract equipment mentions from text using keyword list + model token patterns."""
    if not isinstance(text, str) or not text.strip():
        return []
    found: Set[str] = set()
    low = " " + text.lower() + " "

    # ✅ IMPORTANT: correct regex boundaries (no HTML encoding)
    for kw in keywords:
        kn = kw.lower().strip()
        if not kn:
            continue
        patt = r"(?<!\w)" + re.escape(kn) + r"(?!\w)"
        if re.search(patt, low):
            found.add(kw)

    for pat in MODEL_TOKEN_PATTERNS:
        for m in re.finditer(pat, text):
            found.add(m.group(0))

    return sorted(found, key=lambda s: s.lower())

def equipment_exists_in_doc(equipment: str, doc_text: str) -> bool:
    """Case-insensitive substring check in the Word doc text."""
    eq = re.sub(r"\s+", " ", str(equipment)).strip().lower()
    if not eq:
        return False
    return eq in doc_text


# ---------------------------
# Load testcases (Excel/CSV) with header handling
# ---------------------------
def _coerce_headers(df: pd.DataFrame) -> pd.DataFrame:
    df.columns = [re.sub(r"\s+", " ", str(c).strip()) if not pd.isna(c) else "" for c in df.columns]
    return df

def _norm_col(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (s or "").strip().lower())


def detect_header_row(path: str, sheet: Optional[str], engine: str, env_hint: str) -> int:
    # IMPORTANT: pandas returns dict if sheet_name=None -> use first sheet (0)
    sheet_to_use = 0 if sheet is None else sheet

    tmp = pd.read_excel(
        path,
        sheet_name=sheet_to_use,
        engine=engine,
        header=None,
        nrows=50
    )

    hint = _norm_col(env_hint)
    best = None

    for i in range(min(len(tmp), 50)):
        row_vals = [str(v) if not pd.isna(v) else "" for v in list(tmp.iloc[i])]
        row_norms = [_norm_col(v) for v in row_vals]

        if not any(row_norms):
            continue

        if any(hint in v for v in row_norms):
            return i

        if best is None:
            best = i

    return best if best is not None else 0

def load_testcases(path: str, sheet: Optional[str], header_row: Optional[int], env_hint: str) -> pd.DataFrame:
    if not os.path.exists(path):
        raise FileNotFoundError(f"Testcases file not found: {path}")
    ext = os.path.splitext(path)[1].lower()

    if ext == ".csv":
        df = pd.read_csv(path, header=header_row if header_row is not None else 0)
        return _coerce_headers(df)

    if ext in (".xlsx", ".xls"):
        engine = "openpyxl" if ext == ".xlsx" else "xlrd"

    # IMPORTANT: if sheet is None, read the first sheet (0)
    sheet_to_use = 0 if sheet is None else sheet

    if header_row is None:
        header_row = detect_header_row(path, sheet_to_use, engine, env_hint)

    df = pd.read_excel(path, sheet_name=sheet_to_use, engine=engine, header=header_row)
    return _coerce_headers(df)

    raise ValueError("Unsupported file type. Use .csv, .xlsx, or .xls")

def resolve_env_column(df: pd.DataFrame, desired: str) -> str:
    if desired in df.columns:
        return desired
    target = _norm_col(desired)
    for c in df.columns:
        if _norm_col(c) == target:
            return c

    candidates = []
    for c in df.columns:
        n = _norm_col(c)
        if "environment" in n or "equipment" in n or n.startswith("env"):
            candidates.append(c)

    if len(candidates) == 1:
        print(f"[INFO] Using detected Environment column: {candidates[0]}")
        return candidates[0]
    if len(candidates) > 1:
        print(f"[WARN] Multiple env-like columns found: {candidates}. Using first: {candidates[0]}")
        return candidates[0]

    raise KeyError(f"Environment column '{desired}' not found. Columns: {list(df.columns)}")


# ---------------------------
# Review logic (cell-level per row)
# ---------------------------
def review_env_cells(
    df: pd.DataFrame,
    env_col: str,
    doc_text: str,
    section_numbers: set[str],
    equip_keywords: List[str],
    header_row: Optional[int] = None
) -> Tuple[pd.DataFrame, dict, pd.DataFrame]:
    out = df.copy()

    out["Env_Empty_Flag"] = False
    out["Env_Has_Reference_Flag"] = False
    out["Env_Refs_List"] = ""
    out["Env_Refs_All_Found_Flag"] = True
    out["Env_Refs_Missing_List"] = ""

    out["Env_Instructions_Flag"] = False

    out["Env_Equipment_List"] = ""
    out["Env_Equipment_Flag"] = False
    out["Env_Equipment_All_In_Doc_Flag"] = True
    out["Env_Equipment_Not_In_Doc_List"] = ""

    out["Env_EquipOrInstr_Required_Flag"] = True
    out["Env_Cell_Result"] = ""

    empty_indices = []

    for idx, row in out.iterrows():
        raw = row.get(env_col, None)

        # 1) Empty check
        empty = is_empty_cell(raw)
        out.at[idx, "Env_Empty_Flag"] = empty
        cell_text = "" if empty else str(raw)

        # 2) Reference check (uses section_numbers for auto headings)
        refs = extract_section_refs(cell_text)
        has_ref = len(refs) > 0
        out.at[idx, "Env_Has_Reference_Flag"] = has_ref
        out.at[idx, "Env_Refs_List"] = "; ".join(refs)

        missing_refs = []
        if has_ref:
            for r in refs:
                if not section_exists_in_doc(r, doc_text, section_numbers):
                    missing_refs.append(r)

        out.at[idx, "Env_Refs_Missing_List"] = "; ".join(missing_refs)
        out.at[idx, "Env_Refs_All_Found_Flag"] = (len(missing_refs) == 0)

        # 3) Instructions + equipment detection
        instr = looks_like_instructions(cell_text)
        equip_list = extract_equipment(cell_text, equip_keywords)
        equip_found = len(equip_list) > 0

        out.at[idx, "Env_Instructions_Flag"] = instr
        out.at[idx, "Env_Equipment_Flag"] = equip_found
        out.at[idx, "Env_Equipment_List"] = "; ".join(equip_list)

        # Equipment must exist in docx
        equip_not_in_doc = []
        if equip_found:
            for eq in equip_list:
                if not equipment_exists_in_doc(eq, doc_text):
                    equip_not_in_doc.append(eq)

        out.at[idx, "Env_Equipment_Not_In_Doc_List"] = "; ".join(equip_not_in_doc)
        out.at[idx, "Env_Equipment_All_In_Doc_Flag"] = (len(equip_not_in_doc) == 0)

        # Rule: if not empty and not reference => must have equipment OR instructions
        required_ok = True
        if (not empty) and (not has_ref):
            required_ok = equip_found or instr
        out.at[idx, "Env_EquipOrInstr_Required_Flag"] = required_ok

        # Classification
        if empty:
            out.at[idx, "Env_Cell_Result"] = "EMPTY_ENV"
            empty_indices.append(idx)
        elif has_ref:
            out.at[idx, "Env_Cell_Result"] = "HAS_REFERENCE" if len(missing_refs) == 0 else "WRONG_REFERENCE"
        else:
            if not required_ok:
                out.at[idx, "Env_Cell_Result"] = "MISSING_EQUIPMENT_OR_INSTRUCTIONS"
            elif equip_found and len(equip_not_in_doc) > 0:
                out.at[idx, "Env_Cell_Result"] = "EQUIPMENT_NOT_IN_DOC"
            else:
                out.at[idx, "Env_Cell_Result"] = "OK"

    # ✅ Fixed summary expression (clean and correct)
    summary = {
        "Total rows": len(out),
        "Number of empty rows (Env cell)": int(out["Env_Empty_Flag"].sum()),
        "Rows with references in Env cell": int(out["Env_Has_Reference_Flag"].sum()),
        "Rows with wrong/missing references": int((out["Env_Has_Reference_Flag"] & ~out["Env_Refs_All_Found_Flag"]).sum()),
        "Rows missing equip/instructions when required": int(
            (~out["Env_Empty_Flag"] & ~out["Env_Has_Reference_Flag"] & ~out["Env_EquipOrInstr_Required_Flag"]).sum()
        ),
        "Rows where equipment mentioned but NOT found in doc": int((out["Env_Equipment_Flag"] & ~out["Env_Equipment_All_In_Doc_Flag"]).sum()),
        #"Auto-numbered section labels extracted from DOCX": len(section_numbers),
    }

    # Better Excel row estimate using header_row if provided:
    # Excel row number ≈ header_row (0-based) + 1(header line) + 1(for 1-based rows) + idx
    # => idx + header_row + 2
    if header_row is None:
        excel_rows = [i + 2 for i in empty_indices]  # fallback approximation
    else:
        excel_rows = [i + header_row + 2 for i in empty_indices]

    df_empty_rows = pd.DataFrame({
        "DataFrame_Row_Index": empty_indices,
        "Excel_Row_Number_Estimate": excel_rows,
    })

    return out, summary, df_empty_rows


# ---------------------------
# Safe report writing (avoids PermissionError/locks)
# ---------------------------
def _unique_path(path: str) -> str:
    """If path exists, append timestamp to avoid overwrite/lock issues."""
    if not os.path.exists(path):
        return path
    ts = time.strftime("%Y%m%d_%H%M%S")
    base, ext = os.path.splitext(path)
    return f"{base}_{ts}{ext}"

def save_report(df_review: pd.DataFrame, summary: dict, df_empty: pd.DataFrame, out_path: str) -> str:
    out_path = os.path.abspath(out_path)
    out_dir = os.path.dirname(out_path)
    os.makedirs(out_dir, exist_ok=True)

    base, ext = os.path.splitext(out_path)
    ext = ext.lower()
    if ext not in (".xlsx", ".csv"):
        ext = ".xlsx"
        out_path = base + ext
        base = os.path.splitext(out_path)[0]

    # CSV mode (write unique set)
    if ext == ".csv":
        csv_path = _unique_path(out_path)
        base_csv = os.path.splitext(csv_path)[0]
        df_review.to_csv(csv_path, index=False)
        pd.DataFrame(list(summary.items()), columns=["Metric", "Value"]).to_csv(_unique_path(base_csv + "_summary.csv"), index=False)
        df_empty.to_csv(_unique_path(base_csv + "_empty_rows.csv"), index=False)
        return csv_path

    # XLSX mode: write temp then move to final unique path
    final_path = _unique_path(out_path)

    fd, tmp_name = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)

    try:
        with pd.ExcelWriter(tmp_name, engine="openpyxl") as w:
            df_review.to_excel(w, index=False, sheet_name="Review")
            pd.DataFrame(list(summary.items()), columns=["Metric", "Value"]).to_excel(w, index=False, sheet_name="Summary")
            df_empty.to_excel(w, index=False, sheet_name="EmptyRows")

        shutil.move(tmp_name, final_path)
        return final_path

    except PermissionError as e:
        # fallback to CSV if destination is locked
        try:
            if os.path.exists(tmp_name):
                os.remove(tmp_name)
        except Exception:
            pass
        print(f"[WARN] Permission denied writing XLSX: {e}. Falling back to CSV.", file=sys.stderr)
        csv_path = _unique_path(base + ".csv")
        base_csv = os.path.splitext(csv_path)[0]
        df_review.to_csv(csv_path, index=False)
        pd.DataFrame(list(summary.items()), columns=["Metric", "Value"]).to_csv(_unique_path(base_csv + "_summary.csv"), index=False)
        df_empty.to_csv(_unique_path(base_csv + "_empty_rows.csv"), index=False)
        return csv_path


# ---------------------------
